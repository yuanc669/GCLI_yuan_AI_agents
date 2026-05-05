import os
from pptx import Presentation
from pptx.util import Inches, Pt

def add_slide(prs, title, images, layout_type=1):
    """
    images: list of (path, left, top, width, height)
    """
    slide = prs.slides.add_slide(prs.slide_layouts[layout_type])
    slide.shapes.title.text = title
    
    for img_path, left, top, width, height in images:
        if os.path.exists(img_path):
            slide.shapes.add_picture(img_path, Inches(left), Inches(top), width=Inches(width) if width else None, height=Inches(height) if height else None)
        else:
            print(f"Warning: Image not found: {img_path}")
    return slide

def create_ppt():
    prs = Presentation()
    
    # Root directory for images
    base_dir = r"100_Research\02_Active\DN_GaExo\02_Analysis"
    
    # Slide 1: Title
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "DN_GaExo: Path 1~7 腸道菌叢與生理數據整合分析"
    slide.placeholders[1].text = "大蒜外泌體治療糖尿病腎病變 (DN) 專案\nPath Analysis Visual Summary Report\n2026-05-05"

    # Path 1: Aging Baseline
    imgs_p1 = [
        (os.path.join(base_dir, r"06_Prism_Outputs\Path1_Aging\Path1_Aging_Biochem_Summary.png"), 0.5, 1.5, 4.5, None),
        (os.path.join(base_dir, r"01_Microbiome\02_Diversity\20260502_Aging_Alpha_Plot.png"), 5.2, 1.5, 2.2, None),
        (os.path.join(base_dir, r"01_Microbiome\02_Diversity\20260502_Aging_Beta_PCoA_Plot.png"), 7.5, 1.5, 2.2, None),
        (os.path.join(base_dir, r"01_Microbiome\03_Taxonomy\20260502_Path1_Stacked_Bar_Plot.png"), 0.5, 4.2, 4.0, None),
        (os.path.join(base_dir, r"01_Microbiome\04_Correlation\20260502_Path1_Spearman_Heatmap.png"), 4.8, 4.2, 2.5, None),
        (os.path.join(base_dir, r"01_Microbiome\04_Correlation\20260502_Path1_Venn_Diagram.png"), 7.5, 4.2, 2.0, None)
    ]
    add_slide(prs, "Path 1: 成年期自然老化基準線 (Aging Baseline)", imgs_p1)

    # Path 2: Acute Kidney Injury
    imgs_p2 = [
        (os.path.join(base_dir, r"06_Prism_Outputs\Path2_Acute\Path2_Acute_Biochem_Summary.png"), 0.5, 1.5, 5.0, None),
        (os.path.join(base_dir, r"01_Microbiome\02_Diversity\20260502_Acute_Alpha_Plot.png"), 5.8, 1.5, 2.0, None),
        (os.path.join(base_dir, r"01_Microbiome\02_Diversity\20260502_Acute_Beta_PCoA_Plot.png"), 7.8, 1.5, 2.0, None),
        (os.path.join(base_dir, r"01_Microbiome\03_Taxonomy\20260502_Path2_Stacked_Bar_Plot.png"), 0.5, 4.5, 4.5, None),
        (os.path.join(base_dir, r"01_Microbiome\04_Correlation\20260502_Path2_Spearman_Heatmap.png"), 5.5, 4.5, 3.5, None)
    ]
    add_slide(prs, "Path 2: 急性腎損傷期 (Acute Kidney Injury)", imgs_p2)

    # Path 3: Late/Chronic Stage
    imgs_p3 = [
        (os.path.join(base_dir, r"06_Prism_Outputs\Path3_Late\Path3_Late_Biochem_Summary.png"), 0.5, 1.5, 5.0, None),
        (os.path.join(base_dir, r"01_Microbiome\02_Diversity\20260502_Chronic_Alpha_Plot.png"), 5.8, 1.5, 2.0, None),
        (os.path.join(base_dir, r"01_Microbiome\02_Diversity\20260502_Chronic_Beta_PCoA_Plot.png"), 7.8, 1.5, 2.0, None),
        (os.path.join(base_dir, r"01_Microbiome\03_Taxonomy\20260502_Path3_Stacked_Bar_Plot.png"), 0.5, 4.5, 4.5, None),
        (os.path.join(base_dir, r"01_Microbiome\04_Correlation\20260502_Path3_Spearman_Heatmap.png"), 5.5, 4.5, 3.5, None)
    ]
    add_slide(prs, "Path 3: 慢性/晚期病程 (Late/Chronic Stage)", imgs_p3)

    # Path 4: Disease Progression
    imgs_p4 = [
        (os.path.join(base_dir, r"06_Prism_Outputs\Path4_Progression\Path4_Progression_Biochem_Summary.png"), 0.5, 1.5, 5.0, None),
        (os.path.join(base_dir, r"01_Microbiome\02_Diversity\20260502_Progression_Alpha_Plot.png"), 5.8, 1.5, 2.0, None),
        (os.path.join(base_dir, r"01_Microbiome\02_Diversity\20260502_Progression_Beta_PCoA_Plot.png"), 7.8, 1.5, 2.0, None),
        (os.path.join(base_dir, r"01_Microbiome\04_Correlation\20260502_Path4_Progression_Heatmap.png"), 0.5, 4.5, 4.5, None),
        (os.path.join(base_dir, r"01_Microbiome\03_Taxonomy\20260502_Path4_Progression_Stacked_Bar.png"), 5.5, 4.5, 4.0, None)
    ]
    add_slide(prs, "Path 4: 病程進展動態 (Disease Progression)", imgs_p4)

    # Path 5: Model Drivers
    imgs_p5 = [
        (os.path.join(base_dir, r"06_Prism_Outputs\Path5_Drivers\Path5_Drivers_Biochem_Summary.png"), 0.5, 1.5, 5.0, None),
        (os.path.join(base_dir, r"01_Microbiome\02_Diversity\20260502_Drivers_Alpha_Plot.png"), 5.8, 1.5, 2.0, None),
        (os.path.join(base_dir, r"01_Microbiome\02_Diversity\20260502_Drivers_Beta_PCoA_Plot.png"), 7.8, 1.5, 2.0, None),
        (os.path.join(base_dir, r"01_Microbiome\04_Correlation\20260502_Path5_Driver_Comparison_Plot.png"), 2.5, 4.5, 5.0, None)
    ]
    add_slide(prs, "Path 5: 不同疾病模型驅動因素 (Model Drivers)", imgs_p5)

    # Path 6: Therapy Efficacy
    imgs_p6 = [
        (os.path.join(base_dir, r"06_Prism_Outputs\Path6_Therapy\Path6_Therapy_Biochem_Summary.png"), 0.5, 1.5, 5.0, None),
        (os.path.join(base_dir, r"01_Microbiome\02_Diversity\20260502_Therapy_Alpha_Plot.png"), 5.8, 1.5, 2.0, None),
        (os.path.join(base_dir, r"01_Microbiome\02_Diversity\20260502_Therapy_Beta_PCoA_Plot.png"), 7.8, 1.5, 2.0, None),
        (os.path.join(base_dir, r"01_Microbiome\04_Correlation\20260502_Path6_Therapy_Comparison_Plot.png"), 0.5, 4.5, 4.5, None),
        (os.path.join(base_dir, r"01_Microbiome\04_Correlation\20260502_Path6_Therapy_Heatmap.png"), 5.5, 4.5, 4.0, None)
    ]
    add_slide(prs, "Path 6: 大蒜外泌體治療效果 (Therapy Efficacy)", imgs_p6)

    # Path 7: Global Synthesis
    imgs_p7 = [
        (os.path.join(base_dir, r"01_Microbiome\02_Diversity\20260502_Full_Alpha_Plot.png"), 0.5, 1.5, 4.5, None),
        (os.path.join(base_dir, r"01_Microbiome\02_Diversity\20260502_Full_Beta_PCoA_Plot.png"), 5.2, 1.5, 4.5, None),
        (os.path.join(base_dir, r"01_Microbiome\04_Correlation\20260502_Path7_Global_Heatmap.png"), 2.0, 4.5, 6.0, None)
    ]
    add_slide(prs, "Path 7: 全域整合分析 (Global Synthesis)", imgs_p7)

    output_path = r"100_Research\02_Active\DN_GaExo\02_Analysis\05_Integrated_Synthesis\20260505_DN_GaExo_Visual_Summary.pptx"
    prs.save(output_path)
    print(f"Presentation saved to: {output_path}")

if __name__ == "__main__":
    create_ppt()
